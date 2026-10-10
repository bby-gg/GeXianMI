#!/usr/bin/env python
# encoding: utf-8
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2)
    s.left_margin = Cm(2.5); s.right_margin = Cm(2)

st = doc.styles['Normal']
st.paragraph_format.line_spacing = Pt(16)
st.paragraph_format.space_before = Pt(0); st.paragraph_format.space_after = Pt(0)
st.font.name = '宋体'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def P(text, bold=False, center=False, indent=True, size=12):
    par = doc.add_paragraph()
    par.paragraph_format.line_spacing = Pt(16)
    par.paragraph_format.space_before = Pt(0); par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    if center: par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run(text)
    r.font.name = '宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    r.font.size = Pt(size); r.bold = bold

def H(text, level=1):
    h = doc.add_heading(text, level=level)
    if h.runs:
        h.runs[0].font.name = '黑体'
        h.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        h.runs[0].font.color.rgb = RGBColor(0,0,0)
        h.runs[0].font.size = Pt({1:14,2:13,3:12}[level])

def cap(text):
    par = doc.add_paragraph()
    par.paragraph_format.first_line_indent = Cm(0)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run(text); r.font.size = Pt(10); r.bold = True
    r.font.name = '宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def img(text):
    par = doc.add_paragraph()
    par.paragraph_format.first_line_indent = Cm(0)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run('[' + text + ']')
    r.font.size = Pt(10); r.italic = True
    r.font.color.rgb = RGBColor(128,128,128)
    r.font.name = '宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def tbl(headers, data):
    t = doc.add_table(rows=1+len(data), cols=len(headers), style='Table Grid')
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ''
        r = c.paragraphs[0].add_run(h)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r.bold = True; r.font.size = Pt(10)
        r.font.name = '宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            c = t.rows[ri+1].cells[ci]; c.text = ''
            r = c.paragraphs[0].add_run(str(val))
            r.font.size = Pt(10)
            r.font.name = '宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    doc.add_paragraph()

def BR(): doc.add_page_break()

# ======== 封面 ========
for _ in range(3): doc.add_paragraph()
P('宜春学院人工智能与信息工程学院实训', bold=True, center=True, size=16, indent=False)
doc.add_paragraph(); doc.add_paragraph()
P('基于Spring Boot的葛仙米生长环境', bold=True, center=True, size=20, indent=False)
P('智能监测系统的设计与实现', bold=True, center=True, size=20, indent=False)
for _ in range(4): doc.add_paragraph()
P('学    院：人工智能与信息工程学院', center=True, size=14, indent=False)
P('专    业：计算机科学与技术', center=True, size=14, indent=False)
P('姓    名：___________', center=True, size=14, indent=False)
P('学    号：___________', center=True, size=14, indent=False)
P('指导教师：___________', center=True, size=14, indent=False)
P('日    期：2026年7月', center=True, size=14, indent=False)
doc.add_paragraph()
P('宜春学院人工智能与信息工程学院教务处制', center=True, size=12, indent=False)
BR()

# ======== 摘要 ========
H('基于Spring Boot的葛仙米生长环境智能监测系统的设计与实现', 1)
doc.add_paragraph()
P('摘要：', bold=True, indent=False)
P('本文介绍了葛仙米生长环境智能监测系统的研究背景、项目意义和目前的研究与应用现状，明确了研究内容和主要工作。在业务分析中，对系统存在的问题、原有业务流程等进行了分析，给出了优化后的业务流程。在功能分析中，通过用例分析和用例描述，明确了系统功能。在数据分析中，给出了实体类关系图，建立了数据库表结构。系统采用前后端分离的B/S架构，后端基于Spring Boot + MyBatis-Plus，前端采用Vue 3 + Element Plus + ECharts，涵盖实时监测、历史分析、大棚设备管理、告警监控与规则配置、用户认证等功能模块。')
doc.add_paragraph()
P('关键词：葛仙米；环境监测系统；Spring Boot；Vue 3；需求分析；数据库建模', bold=True, indent=False)
doc.add_paragraph()
P('Abstract：This thesis introduces the Gexianmi Growth Environment Intelligent Monitoring System. The system adopts a B/S architecture with Spring Boot backend and Vue 3 frontend, covering real-time monitoring, historical analysis, greenhouse and device management, alert monitoring and threshold configuration.', bold=False, indent=False)
doc.add_paragraph()
P('Keywords：Gexianmi; Environmental Monitoring; Spring Boot; Vue 3; Requirement Analysis', bold=True, indent=False)
BR()

# ======== 1 ========
H('1 引言', 1)
H('1.1 编写目的', 2)
P('本文编写目的在于规定和控制项目开发内容，明确系统应达到的目标，对系统目标做出完整准确的要求，保证项目需求分析活动在受控状态下进行。')
H('1.2 读者对象', 2)
P('本文档面向项目开发团队、测试人员、项目管理人员及指导教师。')
H('1.3 术语解释', 2)
cap('表1-1 术语解释表')
tbl(['术语','全称','说明'],
    [['B/S','Browser/Server','浏览器/服务器架构'],
     ['JWT','JSON Web Token','身份认证加密令牌'],
     ['ORM','Object Relational Mapping','对象关系映射'],
     ['ECharts','Enterprise Charts','JavaScript图表库'],
     ['BCrypt','—','密码哈希加密函数'],
     ['CORS','Cross-Origin Resource Sharing','跨域资源共享']])
BR()

# ======== 2 ========
H('2 项目概述', 1)
H('2.1 项目背景', 2)
P('葛仙米是一种高经济价值的水生食用藻类，养殖过程对水温、pH值、溶解氧、光照、电导率等环境参数要求严格。传统管理依赖人工定时测量和纸质记录，存在数据采集不及时、数据分散难追溯、异常无法自动告警等问题。随着物联网传感器技术的发展和Web技术的成熟，构建低成本易部署的环境监测系统成为可行方案。本系统旨在实现集数据采集、存储、展示、告警于一体的智能环境监测平台。')
H('2.2 项目目标', 2)
P('(1) 构建Web端环境监测大屏，实时展示5项环境参数的仪表盘和走势曲线。')
P('(2) 实现历史数据查询、可视化分析和CSV数据导出。')
P('(3) 实现大棚和设备的统一管理，自动统计棚区设备数量。')
P('(4) 建立告警管理机制，支持阈值在线配置和告警记录处理。')
P('(5) 实现基于JWT的用户认证和权限控制。')
BR()

# ======== 3 ========
H('3 系统总体功能', 1)
P('系统分为七个核心功能模块：登录注册、实时监测、历史数据、大棚管理、设备管理、告警记录和告警规则配置。')
img('图3-1 系统功能结构图。用层次结构图展示7个功能模块及其子功能')
P('登录注册模块是整个系统的入口，基于JWT实现无状态认证。实时监测模块是操作员日常使用频率最高的功能，展示5参数仪表盘、1小时迷你走势曲线、设备状态表和活跃告警信息。历史数据模块提供时间范围查询和CSV数据导出。大棚管理模块与设备管理模块通过外键关联，实现棚区与设备的联动管理。告警记录模块与告警规则模块协同工作，规则模块配置阈值，记录模块存储告警历史。')
BR()

# ======== 4 ========
H('4 业务需求分析', 1)
H('4.1 环境监测业务', 2)
P('葛仙米养殖环境监测主要涉及五项业务：(1)环境数据采集——目前人工定时测量，约15-20分钟/次，记录分散在纸质或Excel中；(2)数据监控——靠人工目视检查仪表，发现问题后电话通知，响应速度慢；(3)设备管理——设备信息记录在纸质台账，新增更换维护缺乏系统化；(4)告警处理——完全依赖人工判断，缺乏自动告警机制；(5)大棚管理——多个大棚区域需统一管理维护。')
img('图4-1 原有环境监测业务活动图。泳道图展示人工测量记录汇总审核全流程')
P('引入本系统后的优化流程：传感器自动采集数据→通过API上报至服务器→Web大屏实时展示仪表盘和走势曲线→系统自动检测异常并生成告警→操作员即时查看确认。实现了三大改进：数据采集从人工变为全自动、告警从被动发现变为主动推送、数据追溯从翻找纸质记录变为一键查询。')
img('图4-2 优化后环境监测业务活动图。展示采集上报存储展示告警确认的优化流程')
H('4.2 大棚与设备管理业务', 2)
P('管理员需对大棚信息和设备信息进行日常维护。新增养殖筒时需注册设备并绑定到对应大棚。优化后流程：管理员登录→进入管理页→CRUD操作→系统自动维护关联关系→列表自动统计各棚设备数量。')
img('图4-3 大棚设备管理业务活动图')
H('4.3 告警处理业务', 2)
P('传感器数据超出阈值时系统自动生成告警记录并展示。操作员查看详情后可确认（已知晓处理）或忽略（误报）。管理员可在告警规则配置页在线调整各参数阈值。')
img('图4-4 告警处理业务活动图')
BR()

# ======== 5 ========
H('5 系统功能需求', 1)
H('5.1 系统总用例图', 2)
img('图5-1 系统总用例图。三个参与者与7个用例包的关联关系')
H('5.1.1 角色分析', 3)
tbl(['角色','职责','涉及功能'],
    [['管理员','系统最高权限。管理大棚设备、配置告警规则、管理用户','全部7个模块'],
     ['操作员','查看实时和历史数据、确认处理告警','监测/历史/告警'],
     ['观察者','仅查看数据，不修改配置','查看类模块']])
H('5.2 详细功能分析', 2)
cap('表5-1 登录用例描述')
tbl(['描述项','描述内容'],
    [['用例名称','用户登录'],
     ['参与者','管理员、操作员、观察者'],
     ['前置条件','用户已注册，拥有有效用户名密码'],
     ['基本事件流','(1)访问系统显示登录页 (2)输入用户名密码 (3)点击登录 (4)POST /api/auth/login (5)BCrypt验证密码 (6)签发JWT Token(24h有效) (7)返回token和userInfo (8)前端存localStorage (9)跳转/monitor'],
     ['异常事件流','用户名或密码错误提示；Token过期返回401清除Token跳回登录页'],
     ['后置条件','axios拦截器自动注入Authorization请求头'],
     ['优先级','高(1级)']])
cap('表5-2 实时监测用例描述')
tbl(['描述项','描述内容'],
    [['用例名称','实时环境数据监测'],
     ['参与者','管理员、操作员、观察者'],
     ['基本事件流','(1)默认选第一个大棚和设备 (2)fetchLatest()获取最新数据 (3)fetchRecent()获取最近60条用于走势 (4)fetchActiveAlerts()获取告警 (5)渲染5个仪表盘(绿/橙/红三段) (6)渲染1h迷你走势 (7)渲染设备表和告警条 (8)每5秒重复'],
     ['异常事件流','无数据显示等待数据；夜间光照自动切换夜间模式(0-200lux)']])
cap('表5-3 历史数据与大棚管理用例描述')
tbl(['用例','描述'],
    [['历史数据查询','(1)选设备后加载数据范围 (2)快捷时间按钮或手动选时 (3)fetchHistory()查询 (4)双曲线图渲染 (5)可选导出CSV(UTF-8 BOM)'],
     ['大棚信息管理','(1)进入页面加载列表含设备数 (2)支持搜索过滤 (3)新增编辑通过对话框 (4)保存刷新并提示 (5)删除时检查设备数非0阻止']])
cap('表5-4 告警规则配置用例描述')
tbl(['描述项','描述内容'],
    [['用例名称','告警规则配置'],
     ['参与者','管理员'],
     ['基本事件流','(1)进入页面加载5条规则 (2)表格展示参数单位阈值状态 (3)点编辑弹出对话框 (4)修改正常范围告警阈值告警级别启用开关 (5)保存PUT /api/alert-rules/{id}仪表盘立即生效']])
BR()

# ======== 6 ========
H('6 系统需求优先级', 1)
cap('表6-1 系统需求优先级表')
tbl(['编号','功能','优先级','说明'],
    [['REQ-001','用户登录','1级','系统入口'],
     ['REQ-002','5参数仪表盘','1级','核心展示'],
     ['REQ-003','5秒轮询','1级','实时性保障'],
     ['REQ-004','历史数据查询','1级','核心分析'],
     ['REQ-005','大棚管理CRUD','1级','基础设施'],
     ['REQ-006','设备管理CRUD','1级','基础设施'],
     ['REQ-007','告警确认忽略','1级','异常处理'],
     ['REQ-008','用户注册','2级','便利性'],
     ['REQ-009','CSV导出','2级','数据辅助'],
     ['REQ-010','告警阈值配置','2级','灵活性']])
BR()

# ======== 7 ========
H('7 非功能需求', 1)
H('7.1 性能需求', 2)
P('实时监测接口响应时间应在1秒以内；历史数据查询1000条内应在3秒以内。系统应支持至少5个传感器同时上报，10个用户同时查看。')
H('7.2 安全需求', 2)
P('所有API请求除登录注册外需携带JWT Token。密码BCrypt加密存储。前端路由守卫拦截未登录访问。SQL使用MyBatis-Plus参数化查询防注入。')
H('7.3 运行需求', 2)
P('后端：JDK 17+、Maven 3.9+、MySQL 8.0。前端：Node.js 20.19+、现代浏览器。服务器：CPU 4核+、内存8GB+。')
BR()

# ======== 8 ========
H('8 系统总体设计', 1)
H('8.1 体系结构设计', 2)
P('系统采用B/S架构。B/S结构即浏览器/服务器结构，客户端通过浏览器访问服务器上的应用程序，无需安装客户端软件，更新方便易于集中控制。根据葛仙米养殖基地需求，决定采用B/S架构。系统分为四层：表示层Vue 3+Element Plus+ECharts负责界面展示和交互；控制层Spring MVC @RestController+JWT拦截器负责请求校验和路由；业务层Spring @Service封装核心业务逻辑；数据层MySQL 8.0+HikariCP连接池负责数据持久化。')
img('图8-1 系统体系结构图。四层B/S架构：浏览器端HTTP/JSON SpringBoot服务器(Controller Service Mapper) JDBC MySQL')
H('8.2 系统功能架构设计', 2)
P('系统分为七个功能模块：登录注册、实时监测、历史数据、大棚管理、设备管理、告警记录、告警规则。登录注册为基础模块，所有其他模块依赖其进行认证。实时监测模块展示5参数仪表盘和1h迷你走势。历史数据模块提供曲线查询和CSV导出。大棚和设备管理模块实现CRUD操作。告警记录与告警规则模块协同实现告警的全生命周期管理。')
img('图8-2 系统功能架构设计图。层次结构图展示7大功能模块')
H('8.2.1 实时监测功能设计', 3)
P('实时监测是本系统核心展示功能。用户通过大棚和设备下拉框选择目标后，系统展示5个ECharts仪表盘水温/pH/溶氧/光照/电导率、1小时迷你走势5条折线上下排列只读、设备状态表和活跃告警条。仪表盘弧线采用绿正常/橙警告/红异常三段色。光照仪表盘根据数据时间戳自动切换白天模式6:00-18:00范围0-2000lux和黑夜模式范围0-200lux。数据通过Vue watch监听设备选择变化，每5秒定时轮询最新数据。')
img('图8-3 实时监测功能类图。展示RealTimeMonitor.vue GaugePanel.vue EnvironmentController等类关系')
H('8.2.2 大棚管理功能设计', 3)
P('大棚管理功能采用经典CRUD模式。表格展示编号名称位置面积状态设备数备注创建时间，支持搜索过滤。新增和编辑通过对话框实现。删除时检查棚下设备数不为0则阻止。后端GreenhouseController提供5个API端点，Service层自动统计设备数填充至deviceCount字段。')
img('图8-4 大棚管理功能类图。展示GreenhouseManage.vue GreenhouseController GreenhouseService等类关系')
H('8.3 数据库设计', 2)
P('系统涉及6个实体：用户User、大棚Greenhouse、设备Device、环境数据记录EnvironmentRecord、告警记录Alert和告警规则AlertRule。关系：大棚1:N设备，设备1:N环境记录，设备1:N告警记录，告警规则为独立实体。数据库MySQL 8.0，库名gexianmi，字符集utf8mb4，引擎InnoDB。共6张表：')
img('图8-5 系统实体关系图E-R图。6个实体及其1:N关系')
cap('表8-1 数据库表一览')
tbl(['表名','主键','关键字段','说明'],
    [['users','id','username/role/password_hash','用户认证表'],
     ['greenhouses','id','code/name/location/status/area','大棚表含逻辑删除is_deleted和自动填充created_at/updated_at'],
     ['devices','code(VARCHAR)','name/greenhouse_id(FK)/location/status','设备表外键关联greenhouses.id'],
     ['environment_records','id','device_id(FK)/recorded_at/water_temp/ph/dissolved_oxygen/light_intensity/conductivity','数据量最大的表。索引idx_device_time(device_id,recorded_at)'],
     ['alerts','id','device_id(FK)/param/value/threshold_low/threshold_high/level/status/triggered_at','告警表状态为active/confirmed/ignored'],
     ['alert_rules','id','param(UNIQUE)/param_name/unit/min_val/max_val/threshold_low/threshold_high/level/enabled','5条预设规则独立表']])
H('8.4 系统接口设计', 2)
cap('表8-2 系统接口清单节选')
tbl(['模块','方法','路径','认证','说明'],
    [['认证','POST','/api/auth/login','否','登录返回JWT Token'],
     ['认证','POST','/api/auth/register','否','注册默认operator角色'],
     ['环境','POST','/api/environment/data','是','传感器上报数据'],
     ['环境','GET','/api/environment/latest','是','最新一条数据'],
     ['环境','GET','/api/environment/history','是','时间范围查询'],
     ['环境','GET','/api/environment/recent','是','最近N条迷你走势用'],
     ['大棚','GET POST','/api/greenhouses','是','列表含设备数 新建'],
     ['大棚','PUT DELETE','/api/greenhouses/{id}','是','编辑 删除'],
     ['设备','GET POST','/api/devices','是','列表含大棚名 注册'],
     ['设备','PUT DELETE','/api/devices/{code}','是','编辑 删除'],
     ['告警','GET','/api/alerts /alerts/active','是','条件查询 活跃告警'],
     ['告警','POST','/api/alerts/{id}/confirm /ignore','是','确认 忽略'],
     ['规则','GET PUT','/api/alert-rules','是','列表 更新阈值']])
BR()

# ======== 9-12 ========
H('9 开发环境配置', 1)
tbl(['类别','技术','版本'],
    [['后端框架','Spring Boot','3.4.1'],
     ['ORM','MyBatis-Plus','3.5.7'],
     ['数据库','MySQL','8.0'],
     ['前端框架','Vue 3','3.5.x'],
     ['UI组件库','Element Plus','2.14.x'],
     ['图表库','ECharts','6.1.x'],
     ['构建工具','Vite / Maven','8.0 / 3.9'],
     ['Java','JDK','17 LTS'],
     ['Node.js','Node','20.19+']])
H('10 运行环境配置', 1)
tbl(['类别','要求'],
    [['服务器OS','Windows Server 2019+ / Linux'],
     ['Java','JDK 17+'],
     ['数据库','MySQL 8.0'],
     ['Web容器','Spring Boot内嵌Tomcat端口8000'],
     ['客户端浏览器','Chrome Edge Firefox 90+'],
     ['分辨率','1366x768以上']])
H('11 测试环境', 1)
P('单元测试和集成测试环境与开发环境相同。系统测试和验收测试环境与运行环境相同或更严格。')
H('12 系统出错处理', 1)
P('后端通过GlobalExceptionHandler统一捕获异常，400参数校验业务异常、401未登录、500服务器错误。前端axios拦截器处理401自动清除Token跳转登录页。数据库连接失败时HikariCP自动重试。API请求失败时前端保持当前数据不刷新。')
BR()

# ======== 13-15 ========
H('13 测试基本内容', 1)
cap('表13-1 测试计划表')
tbl(['阶段','内容','负责人','完成标准'],
    [['单元测试','各方法功能验证','开发人员','全部通过'],
     ['集成测试','前后端联调','开发人员','接口正常'],
     ['功能测试','按用例验证','测试人员','通过率大于95%'],
     ['验收测试','按需求验收','项目负责人','覆盖率100%']])
H('14 测试实施计划', 1)
cap('表14-1 接口测试节选')
tbl(['接口','输入','预期结果'],
    [['POST /api/auth/login','admin/admin123','code=200返回token'],
     ['POST /api/auth/login','admin/wrong','code=401'],
     ['GET /api/environment/latest','?deviceId=A-001','code=200'],
     ['GET /api/greenhouses','-','code=200含设备数'],
     ['PUT /api/alert-rules/1','{maxVal:26}','code=200']])
cap('表14-2 功能测试节选')
tbl(['模块','用例','操作','预期结果'],
    [['登录','正确登录','admin/admin123','跳转monitor'],
     ['实时监测','切换设备','选择A-002','仪表盘和走势更新'],
     ['历史数据','24h快捷查询','点击24h按钮','展示数据曲线'],
     ['大棚管理','编辑保存','修改名称后保存','提示大棚信息已更新'],
     ['告警规则','修改阈值','修改水温上限保存','提示规则已更新'],
     ['安全','未登录访问','直接输入/monitor','跳转登录页']])
H('15 测试报告', 1)
cap('表15-1 测试用例汇总表')
tbl(['编号','模块','测试项','操作','预期','实际','状态'],
    [['TC-001','登录','正确登录','admin/admin123','跳转/monitor','','通过'],
     ['TC-002','登录','错误密码','admin/wrong','提示错误','',''],
     ['TC-003','注册','正常注册','newuser/123456','注册成功','',''],
     ['TC-004','监测','默认加载','登录后自动','仪表盘和走势','',''],
     ['TC-005','历史','24h查询','点击24h','曲线展示','',''],
     ['TC-006','历史','CSV导出','点击导出','下载CSV','',''],
     ['TC-007','大棚','新增大棚','完整信息提交','提示已创建','',''],
     ['TC-008','大棚','编辑保存','修改名称','提示已更新','',''],
     ['TC-009','规则','修改阈值','修改上限保存','提示已更新','',''],
     ['TC-010','安全','未登录访问','直接输入monitor','跳转login','',''],
     ['TC-011','告警','确认告警','点击确认','状态变confirmed','',''],
     ['TC-012','告警','批量确认','点击确认全部','全部确认','','']])
BR()

# ======== 16-17 ========
H('16 交付', 1)
tbl(['序号','交付项','格式','说明'],
    [['1','系统源代码','ZIP','java_backend和frontend全部源码'],
     ['2','数据库脚本','SQL+Python','建表和数据模拟脚本'],
     ['3','系统设计说明书','DOCX','本文档'],
     ['4','README文档','MD','前后端技术文档']])
BR()

H('17 总结与展望', 1)
H('17.1 总结', 2)
P('本文围绕葛仙米生长环境智能监测系统的分析与设计，主要完成了以下工作：')
P('(1) 对葛仙米养殖环境监测的业务现状进行了深入调研，识别了传统手工记录方式的问题，提出了基于Web的信息化监测方案。')
P('(2) 采用UML建模技术对系统功能需求进行了分析，识别了三类参与者，绘制了系统总用例图，对核心用例进行了详细描述。')
P('(3) 完成了系统的总体设计，采用B/S四层架构，基于Spring Boot+Vue 3技术栈，设计了7个功能模块，绘制了类图。')
P('(4) 完成了数据库设计，识别了6个核心实体及其关系，设计了6张数据表的完整结构，定义了26个RESTful API端点。')
P('(5) 制定了测试计划并设计了12个测试用例，覆盖功能测试、接口测试和安全测试。')
P('通过本次实训，深入理解了软件工程中需求分析和系统设计的方法论，掌握了UML建模技术的应用，积累了前后端分离架构Web应用的完整设计经验。')
H('17.2 展望', 2)
P('(1) 可扩展移动端App或微信小程序，方便手机查看和接收推送。')
P('(2) 可引入机器学习模型进行预测性告警。')
P('(3) 可扩展图像识别功能自动判断葛仙米生长状态。')
P('(4) 可增加硬件联动实现环境参数自动调控。')
P('(5) 可引入RBAC权限模型支持更细粒度权限管理。')
P('(6) 可增加操作日志功能增强可追溯性。')
BR()

# ======== 参考文献 ========
H('参考文献', 1)
refs = [
    '[1] 齐治昌, 谭庆平, 宁洪. 软件工程(第二版)[M]. 高等教育出版社, 2004.',
    '[2] 王珊, 陈红. 数据库系统原理教程[M]. 清华大学出版社, 1998.',
    '[3] 李刚. 轻量级Java EE企业应用实战(第5版)[M]. 电子工业出版社, 2018.',
    '[4] Spring Boot Reference Documentation[EB/OL]. https://docs.spring.io/spring-boot/docs/current/reference/html/, 2026.',
    '[5] Vue.js官方文档[EB/OL]. https://cn.vuejs.org/, 2026.',
    '[6] MyBatis-Plus官方文档[EB/OL]. https://baomidou.com/, 2026.',
    '[7] Apache ECharts官方文档[EB/OL]. https://echarts.apache.org/, 2026.',
    '[8] Element Plus官方文档[EB/OL]. https://element-plus.org/, 2026.',
]
for ref in refs:
    P(ref, indent=False, size=11)

out = r'C:\Users\bby\Desktop\GXM_DigitalTwin\GeXianMI\实训总结报告-葛仙米-v2.docx'
doc.save(out)
import os
print('OK:', out, f'({os.path.getsize(out)} bytes)')